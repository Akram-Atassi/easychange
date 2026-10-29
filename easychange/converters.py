"""Conversion routes and the backends that perform them.

Each route says "files with extension X can become extension Y using backend B".
A backend is only offered when it is actually usable on this machine (e.g. Word
installed, LibreOffice found, ffmpeg on PATH), so the wheel never shows options
that would fail.
"""
from __future__ import annotations

import csv
import functools
import json
import os
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

IS_WINDOWS = sys.platform == "win32"
NO_WINDOW = 0x08000000 if IS_WINDOWS else 0  # CREATE_NO_WINDOW: no console flash

ALIASES = {"jpeg": "jpg", "tif": "tiff", "htm": "html", "markdown": "md",
           "heif": "heic", "hif": "heic"}

# Order in which targets are shown on the wheel (most useful first).
POPULARITY = [
    "pdf", "docx", "jpg", "png", "txt", "mp3", "mp4", "xlsx", "csv", "pptx",
    "webp", "avif", "html", "odt", "rtf", "doc", "gif", "wav", "json", "md", "ods",
    "odp", "bmp", "tiff", "ico", "flac", "ogg", "m4a", "aac", "opus", "mkv",
    "mov", "avi", "webm", "xls", "ppt",
]


def norm_ext(path_or_ext: str | Path) -> str:
    s = str(path_or_ext)
    ext = (Path(s).suffix or s).lower().lstrip(".")
    return ALIASES.get(ext, ext)


class ConversionError(Exception):
    pass


@dataclass(frozen=True)
class Route:
    src: str
    dst: str
    backend: "Backend"


@dataclass(frozen=True)
class Backend:
    name: str
    available: Callable[[], bool]
    # (src, dst, dst_ext). May return a different path it wrote instead, e.g. a
    # folder of page images for a multi-page PDF.
    convert: Callable[[Path, Path, str], Path | None]


# --------------------------------------------------------------------------- #
# Availability checks (cached: they are called once per wheel opening)
# --------------------------------------------------------------------------- #

def _has_module(name: str) -> bool:
    import importlib.util
    return importlib.util.find_spec(name) is not None


@functools.cache
def _com_app_registered(prog_id: str) -> bool:
    """True if an Office app is registered for COM automation (without launching it)."""
    if not IS_WINDOWS or not _has_module("win32com"):
        return False
    import winreg
    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, prog_id + r"\CLSID"))
        return True
    except OSError:
        return False


@functools.cache
def soffice_path() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    if IS_WINDOWS:
        for base in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")):
            if base:
                p = Path(base) / "LibreOffice" / "program" / "soffice.exe"
                if p.exists():
                    return str(p)
    return None


@functools.cache
def ffmpeg_path() -> str | None:
    found = shutil.which("ffmpeg")
    if found:
        return found
    # A copy of ffmpeg.exe shipped next to EasyChange.exe also works.
    local = Path(sys.executable).parent / "ffmpeg.exe"
    return str(local) if local.exists() else None


# --------------------------------------------------------------------------- #
# Backends
# --------------------------------------------------------------------------- #

def _run(cmd: list[str], timeout: int = 600) -> None:
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout,
                           creationflags=NO_WINDOW)
    except subprocess.TimeoutExpired:
        raise ConversionError("Conversion timed out")
    if r.returncode != 0:
        msg = (r.stderr or r.stdout).strip().splitlines()
        raise ConversionError(msg[-1] if msg else f"{Path(cmd[0]).name} failed")


# ---- Images (Pillow) ------------------------------------------------------ #

@functools.cache
def _heic_ok() -> bool:
    """iPhone photos (HEIC) need the pillow-heif plugin."""
    if not _has_module("pillow_heif"):
        return False
    import pillow_heif
    pillow_heif.register_heif_opener()
    return True


@functools.cache
def _avif_ok() -> bool:
    try:
        from PIL import features
        return bool(features.check_module("avif"))
    except Exception:
        return False


def _pillow_convert(src: Path, dst: Path, ext: str) -> None:
    from PIL import Image, ImageOps, ImageSequence

    if norm_ext(src) == "heic" and not _heic_ok():
        raise ConversionError("HEIC support (pillow-heif) is not installed")
    with Image.open(src) as im:
        if ext == "gif" and getattr(im, "is_animated", False):
            frames = [f.copy() for f in ImageSequence.Iterator(im)]
            frames[0].save(dst, save_all=True, append_images=frames[1:], loop=0)
            return
        im.load()
        # Phone photos are often stored sideways with a "rotate me" tag that many
        # formats can't carry; apply the rotation so the result looks right.
        im = ImageOps.exif_transpose(im)
        if ext in ("jpg", "pdf", "bmp"):
            # Formats without alpha: flatten onto white instead of turning it black.
            if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
                rgba = im.convert("RGBA")
                bg = Image.new("RGB", rgba.size, (255, 255, 255))
                bg.paste(rgba, mask=rgba.getchannel("A"))
                im = bg
            else:
                im = im.convert("RGB")
        fmt = {"jpg": "JPEG", "tiff": "TIFF", "ico": "ICO"}.get(ext, ext.upper())
        opts = {"quality": 92} if ext in ("jpg", "webp") else {}
        if ext == "ico":
            opts["sizes"] = [s for s in [(16, 16), (32, 32), (48, 48), (64, 64),
                                         (128, 128), (256, 256)]
                             if s[0] <= max(im.size)] or [(16, 16)]
        im.save(dst, fmt, **opts)


PILLOW = Backend("Pillow", lambda: _has_module("PIL"), _pillow_convert)
AVIF = Backend("Pillow AVIF", _avif_ok, _pillow_convert)
HEIC = Backend("Pillow HEIC", _heic_ok, _pillow_convert)
HEIC_AVIF = Backend("Pillow HEIC+AVIF", lambda: _heic_ok() and _avif_ok(), _pillow_convert)
IMAGE_IN = ["png", "jpg", "bmp", "gif", "webp", "tiff", "ico"]
IMAGE_OUT = ["png", "jpg", "bmp", "gif", "webp", "tiff", "ico", "pdf"]
IMAGES = IMAGE_IN + ["avif", "heic"]  # everything that counts as an image


# ---- Microsoft Office via COM --------------------------------------------- #

WORD_FORMATS = {"doc": 0, "txt": 7, "rtf": 6, "html": 10, "docx": 16, "pdf": 17, "odt": 23}
EXCEL_FORMATS = {"csv": 6, "xlsx": 51, "xls": 56, "ods": 60}
PPT_FORMATS = {"ppt": 1, "pptx": 24, "pdf": 32, "odp": 35}


def _office(prog_id: str):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx(prog_id)  # a private instance, never the user's open one
    try:
        app.DisplayAlerts = False
    except Exception:
        pass
    return app


def _word_convert(src: Path, dst: Path, ext: str) -> None:
    app = _office("Word.Application")
    try:
        doc = app.Documents.Open(str(src), ReadOnly=True, AddToRecentFiles=False,
                                 ConfirmConversions=False)
        try:
            if ext == "txt":  # plain text, encoded as UTF-8 (code page 65001)
                doc.SaveAs2(str(dst), FileFormat=7, Encoding=65001)
            else:
                doc.SaveAs2(str(dst), FileFormat=WORD_FORMATS[ext])
        finally:
            doc.Close(SaveChanges=0)
    finally:
        app.Quit()


def _excel_convert(src: Path, dst: Path, ext: str) -> None:
    app = _office("Excel.Application")
    try:
        wb = app.Workbooks.Open(str(src), ReadOnly=True, AddToMru=False)
        try:
            if ext == "pdf":
                wb.ExportAsFixedFormat(0, str(dst))
            else:
                wb.SaveAs(str(dst), FileFormat=EXCEL_FORMATS[ext])
        finally:
            wb.Close(SaveChanges=False)
    finally:
        app.Quit()


def _ppt_convert(src: Path, dst: Path, ext: str) -> None:
    app = _office("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(src), ReadOnly=True, WithWindow=False)
        try:
            pres.SaveAs(str(dst), PPT_FORMATS[ext])
        finally:
            pres.Close()
    finally:
        app.Quit()


WORD = Backend("Microsoft Word", lambda: _com_app_registered("Word.Application"), _word_convert)
EXCEL = Backend("Microsoft Excel", lambda: _com_app_registered("Excel.Application"), _excel_convert)
POWERPOINT = Backend("Microsoft PowerPoint",
                     lambda: _com_app_registered("PowerPoint.Application"), _ppt_convert)
EXCEL_FORMATS_OUT = list(EXCEL_FORMATS) + ["pdf"]


# ---- LibreOffice ---------------------------------------------------------- #

LO_FILTERS = {"txt": "txt:Text (encoded):UTF8", "html": "html:XHTML Writer File:UTF8"}


def _lo_convert(src: Path, dst: Path, ext: str) -> None:
    with tempfile.TemporaryDirectory() as out, tempfile.TemporaryDirectory() as profile:
        # A throwaway profile lets this run even while the user has LibreOffice open.
        _run([soffice_path(), f"-env:UserInstallation={Path(profile).as_uri()}",
              "--headless", "--convert-to", LO_FILTERS.get(ext, ext),
              "--outdir", out, str(src)], timeout=300)
        produced = list(Path(out).iterdir())
        if not produced:
            raise ConversionError("LibreOffice could not convert this file")
        shutil.move(str(produced[0]), dst)


LIBREOFFICE = Backend("LibreOffice", lambda: soffice_path() is not None, _lo_convert)
LO_WRITER = ["docx", "doc", "odt", "rtf", "txt", "html", "pdf"]
LO_CALC = ["xlsx", "xls", "ods", "csv", "pdf"]
LO_IMPRESS = ["pptx", "ppt", "odp", "pdf"]


# ---- PDF as a source (pypdfium2: the PDF engine used by Chrome) ---------- #

def _pdf_to_images(src: Path, dst: Path, ext: str) -> Path:
    """One page: report.png. Several pages: a folder 'report (png pages)' with
    'page 01.png', 'page 02.png', ..."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(src)
    try:
        n = len(pdf)
        if n == 0:
            raise ConversionError("This PDF has no pages")
        out = dst
        if n > 1:
            out = src.with_name(f"{src.stem} ({ext} pages)")
            i = 2
            while out.exists():
                out = src.with_name(f"{src.stem} ({ext} pages {i})")
                i += 1
            out.mkdir()
        width = len(str(n))
        try:
            for i in range(n):
                page = pdf[i]
                img = page.render(scale=2, fill_color=(255, 255, 255, 255)).to_pil()  # 144 dpi
                if ext == "jpg":
                    img = img.convert("RGB")
                target = dst if n == 1 else out / f"page {i + 1:0{width}}.{ext}"
                img.save(target, {"jpg": "JPEG"}.get(ext, ext.upper()),
                         **({"quality": 92} if ext in ("jpg", "webp") else {}))
        except Exception:
            if n > 1:
                shutil.rmtree(out, ignore_errors=True)
            raise
        return out
    finally:
        pdf.close()


def _pdf_to_txt(src: Path, dst: Path, ext: str) -> None:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(src)
    try:
        pages = [pdf[i].get_textpage().get_text_range() for i in range(len(pdf))]
    finally:
        pdf.close()
    text = "\n\n".join(p.replace("\r\n", "\n").strip() for p in pages)
    if not text.strip():
        raise ConversionError("No text found (the PDF may be scanned images)")
    dst.write_text(text + "\n", encoding="utf-8")


PDFIUM = Backend("PDF renderer", lambda: _has_module("pypdfium2"), _pdf_to_images)
PDF_TEXT = Backend("PDF text", lambda: _has_module("pypdfium2"), _pdf_to_txt)


# ---- Pure-Python fallbacks (always available) ----------------------------- #

def _docx_to_txt(src: Path, dst: Path, ext: str) -> None:
    import docx
    d = docx.Document(str(src))
    dst.write_text("\n".join(p.text for p in d.paragraphs), encoding="utf-8")


def _read_csv(src: Path) -> list[list[str]]:
    raw = src.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(text.splitlines(), dialect))


def _csv_to_xlsx(src: Path, dst: Path, ext: str) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for row in _read_csv(src):
        ws.append([_maybe_number(c) for c in row])
    wb.save(dst)


def _maybe_number(s: str):
    try:
        return int(s) if s.strip().lstrip("-").isdigit() else float(s) if s.strip() else s
    except ValueError:
        return s


def _xlsx_to_csv(src: Path, dst: Path, ext: str) -> None:
    from openpyxl import load_workbook
    wb = load_workbook(src, read_only=True, data_only=True)
    with dst.open("w", newline="", encoding="utf-8-sig") as f:  # BOM so Excel reads UTF-8
        w = csv.writer(f)
        for row in wb.active.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])


def _csv_to_json(src: Path, dst: Path, ext: str) -> None:
    rows = _read_csv(src)
    if not rows:
        data = []
    else:
        header, body = rows[0], rows[1:]
        data = [dict(zip(header, r)) for r in body]
    dst.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def _json_to_csv(src: Path, dst: Path, ext: str) -> None:
    data = json.loads(src.read_text(encoding="utf-8-sig"))
    if isinstance(data, dict):
        data = [data]
    if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
        raise ConversionError("JSON must be a list of objects to become a table")
    header: list[str] = []
    for r in data:
        header += [k for k in r if k not in header]
    with dst.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=header)
        w.writeheader()
        for r in data:
            w.writerow({k: (json.dumps(v) if isinstance(v, (dict, list)) else v)
                        for k, v in r.items()})


def _md_to_html(src: Path, dst: Path, ext: str) -> None:
    import markdown
    body = markdown.markdown(src.read_text(encoding="utf-8-sig"),
                             extensions=["tables", "fenced_code"])
    dst.write_text(f"<!doctype html>\n<meta charset='utf-8'>\n<title>{src.stem}</title>\n"
                   f"{body}\n", encoding="utf-8")


def _copy_text(src: Path, dst: Path, ext: str) -> None:
    # txt <-> md: the content is already valid in both formats.
    shutil.copyfile(src, dst)


def _py(name: str, fn, *modules: str) -> Backend:
    return Backend(name, lambda: all(_has_module(m) for m in modules), fn)


# ---- Audio / video (ffmpeg) ----------------------------------------------- #

AUDIO = ["mp3", "wav", "flac", "ogg", "m4a", "aac", "opus"]
VIDEO = ["mp4", "mkv", "mov", "avi", "webm"]


def _ffmpeg_convert(src: Path, dst: Path, ext: str) -> None:
    cmd = [ffmpeg_path(), "-hide_banner", "-loglevel", "error", "-y", "-i", str(src)]
    if ext in AUDIO:
        cmd += ["-vn"]
    if ext == "gif":  # palette-based GIF looks far better than the default
        cmd += ["-vf", "fps=12,scale='min(640,iw)':-2:flags=lanczos,split[a][b];"
                       "[a]palettegen[p];[b][p]paletteuse"]
    _run(cmd + [str(dst)], timeout=3600)


FFMPEG = Backend("ffmpeg", lambda: ffmpeg_path() is not None, _ffmpeg_convert)


# --------------------------------------------------------------------------- #
# Route table. Earlier entries win when two backends offer the same pair.
# --------------------------------------------------------------------------- #

def _routes(backend: Backend, srcs, dsts) -> list[Route]:
    return [Route(s, d, backend) for s in srcs for d in dsts if s != d]


ROUTES: list[Route] = [
    # CSV <-> XLSX in Python first: it detects ; and tab separators (Office and
    # LibreOffice assume commas) and writes UTF-8.
    Route("csv", "xlsx", _py("openpyxl", _csv_to_xlsx, "openpyxl")),
    Route("xlsx", "csv", _py("openpyxl", _xlsx_to_csv, "openpyxl")),
    *_routes(PILLOW, IMAGE_IN, IMAGE_OUT),
    *_routes(AVIF, IMAGE_IN, ["avif"]),
    *_routes(AVIF, ["avif"], IMAGE_OUT),
    *_routes(HEIC, ["heic"], IMAGE_OUT),
    Route("heic", "avif", HEIC_AVIF),
    *_routes(WORD, ["docx", "doc", "odt", "rtf"], list(WORD_FORMATS)),
    Route("pdf", "docx", WORD),  # Word turns PDFs into editable documents
    *_routes(PDFIUM, ["pdf"], ["png", "jpg", "webp"]),
    Route("pdf", "txt", PDF_TEXT),
    *_routes(EXCEL, ["xlsx", "xls", "ods", "csv"], EXCEL_FORMATS_OUT),
    *_routes(POWERPOINT, ["pptx", "ppt", "odp"], list(PPT_FORMATS)),
    *_routes(LIBREOFFICE, ["docx", "doc", "odt", "rtf"], LO_WRITER),
    *_routes(LIBREOFFICE, ["txt"], ["docx", "odt", "pdf"]),
    *_routes(LIBREOFFICE, ["xlsx", "xls", "ods", "csv"], LO_CALC),
    *_routes(LIBREOFFICE, ["pptx", "ppt", "odp"], LO_IMPRESS),
    Route("docx", "txt", _py("python-docx", _docx_to_txt, "docx")),
    Route("csv", "json", _py("csv", _csv_to_json)),
    Route("json", "csv", _py("json", _json_to_csv)),
    Route("md", "html", _py("markdown", _md_to_html, "markdown")),
    Route("md", "txt", _py("copy", _copy_text)),
    Route("txt", "md", _py("copy", _copy_text)),
    *_routes(FFMPEG, AUDIO, AUDIO),
    *_routes(FFMPEG, VIDEO, VIDEO + ["gif", "mp3", "wav"]),
]


def _rank(ext: str) -> int:
    return POPULARITY.index(ext) if ext in POPULARITY else len(POPULARITY)


def targets_for(path: str | Path) -> list[str]:
    """Extensions this file can be converted to on this machine, most useful first."""
    src = norm_ext(path)
    seen: list[str] = []
    checked: dict[str, bool] = {}
    for r in ROUTES:
        if r.src != src or r.dst in seen:
            continue
        ok = checked.setdefault(r.backend.name, r.backend.available())
        if ok:
            seen.append(r.dst)
    return sorted(seen, key=_rank)


def unique_output_path(src: Path, ext: str) -> Path:
    """report.docx -> report.pdf, or report (2).pdf if that already exists."""
    out = src.with_suffix("." + ext)
    n = 2
    while out.exists():
        out = src.with_name(f"{src.stem} ({n}).{ext}")
        n += 1
    return out


def _friendly(e: Exception, src: Path) -> str:
    """Error text for the red bar: no full paths, no library jargon."""
    msg = (str(e) or type(e).__name__).replace(f"'{src}'", src.name).replace(str(src), src.name)
    if "cannot identify image file" in msg:
        return "Not a valid image (or the file is damaged)"
    return msg


def convert(path: str | Path, ext: str) -> Path:
    """Convert `path` to `ext` next to the original. The original is never modified."""
    src = Path(path).resolve()
    if not src.is_file():
        raise ConversionError("File not found")
    ext = norm_ext(ext)
    s = norm_ext(src)
    errors = []
    for r in ROUTES:
        if r.src != s or r.dst != ext or not r.backend.available():
            continue
        dst = unique_output_path(src, ext)
        try:
            out = r.backend.convert(src, dst, ext) or dst
        except Exception as e:  # try the next backend that offers this pair
            if dst.is_file():
                dst.unlink(missing_ok=True)
            errors.append(_friendly(e, src))
            continue
        if out.is_dir() or (out.exists() and out.stat().st_size > 0):
            return out
        errors.append(f"{r.backend.name} produced no output")
    if not errors:
        raise ConversionError(f"Can't convert .{s} to .{ext} on this PC")
    raise ConversionError(errors[-1])


# --------------------------------------------------------------------------- #
# Several files at once
# --------------------------------------------------------------------------- #

COMBINE = "combine-pdf"  # special wheel option: merge images/PDFs into one PDF


def label(target: str) -> str:
    """Text for a wheel slot."""
    return "1 PDF" if target == COMBINE else target.upper()


def _can_combine(paths: list[Path]) -> bool:
    if len(paths) < 2 or not _has_module("pypdfium2"):
        return False
    exts = {norm_ext(p) for p in paths}
    if not exts <= set(IMAGES) | {"pdf"}:
        return False
    return ("heic" not in exts or _heic_ok()) and ("avif" not in exts or _avif_ok())


def shared_targets(paths: list[str | Path]) -> list[str]:
    """Targets offered for a selection. With several files, a type is offered if
    every file can become it (files already of that type are just skipped)."""
    paths = [Path(p) for p in paths]
    if len(paths) == 1:
        return targets_for(paths[0])
    per_ext = {}
    for p in paths:
        e = norm_ext(p)
        if e not in per_ext:
            per_ext[e] = set(targets_for(p))
    candidates = set().union(*per_ext.values())
    result = sorted((t for t in candidates
                     if all(t in ts or e == t for e, ts in per_ext.items())), key=_rank)
    if _can_combine(paths):
        result.insert(1 if result[:1] == ["pdf"] else 0, COMBINE)
    return result


def _natural_key(p: Path):
    import re
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", p.name)]


def combine_pdf(paths: list[str | Path]) -> Path:
    """Merge images and PDFs (sorted by name, like Explorer) into 'Combined.pdf'."""
    import io

    import pypdfium2 as pdfium
    from PIL import Image, ImageOps

    files = sorted((Path(p).resolve() for p in paths), key=_natural_key)
    folder = files[0].parent
    dst = folder / "Combined.pdf"
    n = 2
    while dst.exists():
        dst = folder / f"Combined ({n}).pdf"
        n += 1

    out = pdfium.PdfDocument.new()
    sources = []  # keep imported documents open until the result is saved
    try:
        for f in files:
            try:
                if norm_ext(f) == "pdf":
                    doc = pdfium.PdfDocument(f)
                else:
                    if norm_ext(f) == "heic":
                        _heic_ok()
                    with Image.open(f) as im:
                        im = ImageOps.exif_transpose(im)
                        if im.mode != "RGB":
                            rgba = im.convert("RGBA")
                            im = Image.new("RGB", rgba.size, (255, 255, 255))
                            im.paste(rgba, mask=rgba.getchannel("A"))
                        # Size the page so the long side is about 11 inches (A4/Letter).
                        dpi = max(72.0, max(im.size) / 11.0)
                        buf = io.BytesIO()
                        im.save(buf, "PDF", resolution=dpi)
                    doc = pdfium.PdfDocument(buf.getvalue())
            except Exception as e:
                raise ConversionError(f"Couldn't read {f.name} ({e})")
            sources.append(doc)
            out.import_pages(doc)
        out.save(dst)
    finally:
        out.close()
        for d in sources:
            d.close()
    return dst


def convert_many(paths: list[str | Path], target: str, progress=None):
    """Convert a selection. Returns (outputs, failures) where failures is a list of
    (path, message). `progress(done, total)` is called after each file."""
    paths = [Path(p) for p in paths]
    if target == COMBINE:
        return [combine_pdf(paths)], []
    todo = [p for p in paths if norm_ext(p) != target]
    outputs, failures = [], []
    for i, p in enumerate(todo, 1):
        try:
            outputs.append(convert(p, target))
        except Exception as e:
            failures.append((p, str(e) or type(e).__name__))
        if progress:
            progress(i, len(todo))
    return outputs, failures
