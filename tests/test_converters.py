import json
import shutil

import pytest
from PIL import Image

from easychange import converters as c


def test_norm_ext():
    assert c.norm_ext("a/b/Photo.JPEG") == "jpg"
    assert c.norm_ext(".TIF") == "tiff"
    assert c.norm_ext("docx") == "docx"
    assert c.norm_ext("noext") == "noext"


def test_unknown_type_has_no_targets(tmp_path):
    f = tmp_path / "thing.xyz"
    f.write_text("hi")
    assert c.targets_for(f) == []


def test_targets_exclude_self_and_aliases(tmp_path):
    t = c.targets_for(tmp_path / "a.jpeg")
    assert "jpg" not in t and "png" in t and t[0] == "pdf"


def test_unique_output_path(tmp_path):
    src = tmp_path / "r.png"
    src.touch()
    (tmp_path / "r.jpg").touch()
    assert c.unique_output_path(src, "jpg").name == "r (2).jpg"


def test_png_with_alpha_to_jpg_and_pdf(tmp_path):
    src = tmp_path / "logo.png"
    Image.new("RGBA", (40, 30), (255, 0, 0, 0)).save(src)
    out = c.convert(src, "jpg")
    with Image.open(out) as im:
        assert im.format == "JPEG" and im.getpixel((0, 0)) == (255, 255, 255)
    assert c.convert(src, "pdf").read_bytes().startswith(b"%PDF")
    assert c.convert(src, "ico").exists()
    assert src.exists()  # original untouched


def test_csv_json_xlsx_roundtrip(tmp_path):
    src = tmp_path / "t.csv"
    src.write_text("name;age\nAda;36\nLin;7\n", encoding="utf-8")
    j = c.convert(src, "json")
    assert json.loads(j.read_text()) == [{"name": "Ada", "age": "36"}, {"name": "Lin", "age": "7"}]
    x = c.convert(src, "xlsx")
    back = c.convert(x, "csv")
    assert back.name == "t (2).csv"
    assert "Ada,36" in back.read_text(encoding="utf-8-sig")


def test_json_must_be_table(tmp_path):
    src = tmp_path / "x.json"
    src.write_text("[1, 2]")
    with pytest.raises(c.ConversionError, match="list of objects"):
        c.convert(src, "csv")
    assert not (tmp_path / "x.csv").exists()


def test_md_to_html(tmp_path):
    src = tmp_path / "n.md"
    src.write_text("# Hi\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    html = c.convert(src, "html").read_text()
    assert "<h1>Hi</h1>" in html and "<table>" in html


def test_docx_to_txt(tmp_path):
    import docx
    src = tmp_path / "d.docx"
    d = docx.Document()
    d.add_paragraph("Hello wheel")
    d.save(src)
    assert "txt" in c.targets_for(src)
    route = next(r for r in c.ROUTES if r.src == "docx" and r.dst == "txt"
                 and r.backend.name == "python-docx")
    out = tmp_path / "d.txt"
    route.backend.convert(src, out, "txt")
    assert out.read_text().strip() == "Hello wheel"


@pytest.mark.skipif(not c.soffice_path(), reason="LibreOffice not installed")
def test_docx_to_pdf_with_libreoffice(tmp_path):
    import docx
    src = tmp_path / "report.docx"
    d = docx.Document()
    d.add_paragraph("Quarterly report")
    d.save(src)
    assert c.targets_for(src)[0] == "pdf"
    assert c.convert(src, "pdf").read_bytes().startswith(b"%PDF")


@pytest.mark.skipif(not c.ffmpeg_path(), reason="ffmpeg not installed")
def test_wav_to_mp3(tmp_path):
    import subprocess
    src = tmp_path / "beep.wav"
    subprocess.run([c.ffmpeg_path(), "-loglevel", "error", "-f", "lavfi", "-i",
                    "sine=d=0.5", str(src)], check=True)
    assert c.convert(src, "mp3").stat().st_size > 0


def _make_pdf(path, pages):
    colors = ["red", "green", "blue", "white"]
    ims = [Image.new("RGB", (60, 80), colors[i % 4]) for i in range(pages)]
    ims[0].save(path, save_all=True, append_images=ims[1:])


def test_pdf_targets(tmp_path):
    t = c.targets_for(tmp_path / "Scan.PDF")
    assert {"png", "jpg", "webp", "txt"} <= set(t)


def test_single_page_pdf_to_png(tmp_path):
    src = tmp_path / "one.pdf"
    _make_pdf(src, 1)
    out = c.convert(src, "png")
    assert out.name == "one.png"
    with Image.open(out) as im:
        assert im.size == (120, 160)  # rendered at 2x
        r, g, b = im.convert("RGB").getpixel((60, 80))
        assert r > 240 and g < 15 and b < 15  # red (PDF stores it as JPEG)


def test_multi_page_pdf_to_jpg_makes_folder(tmp_path):
    src = tmp_path / "three.pdf"
    _make_pdf(src, 3)
    out = c.convert(src, "jpg")
    assert out.is_dir() and out.name == "three (jpg pages)"
    assert sorted(p.name for p in out.iterdir()) == ["page 1.jpg", "page 2.jpg", "page 3.jpg"]
    assert c.convert(src, "jpg").name == "three (jpg pages 2)"


def test_image_only_pdf_to_txt_explains(tmp_path):
    src = tmp_path / "scan.pdf"
    _make_pdf(src, 1)
    with pytest.raises(c.ConversionError, match="No text found"):
        c.convert(src, "txt")
    assert not (tmp_path / "scan.txt").exists()


@pytest.mark.skipif(not c.soffice_path(), reason="LibreOffice not installed")
def test_pdf_text_roundtrip(tmp_path):
    import docx
    src = tmp_path / "letter.docx"
    d = docx.Document()
    d.add_paragraph("Dear wheel, hello from page one.")
    d.save(src)
    pdf = c.convert(src, "pdf")
    assert "Dear wheel, hello from page one." in c.convert(pdf, "txt").read_text()


@pytest.mark.skipif(not c._heic_ok(), reason="pillow-heif not installed")
def test_heic_to_jpg(tmp_path):
    src = tmp_path / "IMG_0001.HEIC"
    Image.new("RGB", (64, 48), "blue").save(src, "HEIF")
    t = c.targets_for(src)
    assert t[:3] == ["pdf", "jpg", "png"] and "heic" not in t
    with Image.open(c.convert(src, "jpg")) as im:
        assert im.format == "JPEG" and im.size == (64, 48)


@pytest.mark.skipif(not c._avif_ok(), reason="Pillow built without AVIF")
def test_avif_roundtrip(tmp_path):
    src = tmp_path / "a.png"
    Image.new("RGB", (32, 32), "green").save(src)
    assert "avif" in c.targets_for(src)
    avif = c.convert(src, "avif")
    assert "png" in c.targets_for(avif)
    with Image.open(c.convert(avif, "png")) as im:
        assert im.size == (32, 32)


def test_photo_rotation_is_applied(tmp_path):
    src = tmp_path / "phone.jpg"
    exif = Image.Exif()
    exif[0x0112] = 6  # "rotate 90° clockwise to display"
    Image.new("RGB", (40, 20), "red").save(src, exif=exif)
    with Image.open(c.convert(src, "png")) as im:
        assert im.size == (20, 40)


def test_shared_targets_for_a_mixed_selection(tmp_path):
    t = c.shared_targets([tmp_path / "a.png", tmp_path / "b.jpg"])
    # jpg is offered: the png gets converted and the jpg is skipped.
    assert t[:2] == ["pdf", c.COMBINE] and "jpg" in t and "png" in t
    assert c.shared_targets([tmp_path / "a.docx", tmp_path / "b.xyz"]) == []
    assert c.COMBINE not in c.shared_targets([tmp_path / "a.png"])
    assert c.COMBINE not in c.shared_targets([tmp_path / "a.png", tmp_path / "b.docx"])


def test_convert_many_skips_files_already_in_the_target_type(tmp_path):
    a, b, bad = tmp_path / "a.png", tmp_path / "b.jpg", tmp_path / "c.png"
    Image.new("RGB", (8, 8)).save(a)
    Image.new("RGB", (8, 8)).save(b)
    bad.write_bytes(b"not an image")
    seen = []
    outputs, failures = c.convert_many([a, b, bad], "jpg", progress=lambda d, t: seen.append((d, t)))
    assert [o.name for o in outputs] == ["a.jpg"]
    assert [f.name for f, _ in failures] == ["c.png"]
    assert seen == [(1, 2), (2, 2)]
    assert not (tmp_path / "b (2).jpg").exists()


def test_combine_images_and_pdfs_in_name_order(tmp_path):
    import pypdfium2 as pdfium
    Image.new("RGBA", (300, 200), (0, 0, 255, 128)).save(tmp_path / "scan 10.png")
    Image.new("RGB", (200, 300), "red").save(tmp_path / "scan 2.jpg")
    _make_pdf(tmp_path / "scan 1.pdf", 2)
    files = [tmp_path / n for n in ("scan 10.png", "scan 2.jpg", "scan 1.pdf")]
    outputs, failures = c.convert_many(files, c.COMBINE)
    assert failures == [] and outputs[0].name == "Combined.pdf"
    pdf = pdfium.PdfDocument(outputs[0])
    sizes = [pdf[i].get_size() for i in range(len(pdf))]
    pdf.close()
    assert len(sizes) == 4  # 2 pages + 2 images
    assert sizes[2][0] < sizes[2][1] and sizes[3][0] > sizes[3][1]  # "scan 2" before "scan 10"
    assert c.combine_pdf(files).name == "Combined (2).pdf"


def test_combine_reports_unreadable_file(tmp_path):
    good, bad = tmp_path / "a.png", tmp_path / "b.pdf"
    Image.new("RGB", (8, 8)).save(good)
    bad.write_bytes(b"%PDF-broken")
    with pytest.raises(c.ConversionError, match="b.pdf"):
        c.combine_pdf([good, bad])
    assert not (tmp_path / "Combined.pdf").exists()
